"""
Build a standalone knowledge file of stewardship meeting attendees and
strategic timelines, for the Copilot "What's on my plate" agent.

Reads the Cover and Timeline tabs out of every client workbook in the
Stewardships - Documents library (a separate SharePoint library from this
project, so its path is configured below) and writes
Whats_On_My_Plate_Stewardship_Knowledge.md to BOTH this project's
Master-data AND Employee-data folders -- the identical file goes to both,
since stewardship/client data has no Leadership-vs-everyone split the way
Ninety rocks or Monday tasks do, so there's nothing to filter out for the
employee-facing agent. One section per client, in plain sentences so
Copilot retrieval matches natural questions like "who was at Bank of
Utica's last stewardship meeting" or "what's coming up for ARC Otsego".

From the Cover tab: the client name, most recent annual review meeting
date, and attendee list.

From the Timeline tab: every milestone on the client's 450-day strategic
timeline grid (activity name, the calendar date it falls on, the
milestone label, and who owns it -- GKG, the client ("You"), or both),
sorted chronologically.

This script only ever reads files in Stewardships - Documents -- it never
writes to, renames, or deletes anything there.

Run any time you want a fresh copy of this knowledge file:
    python 4-build_stewardship_knowledge_file.py

Not wired into run_daily_pipeline.py yet -- run it manually, or ask to
have it added as a step once you're happy with the output.

PORTABILITY: this script does not hardcode any particular computer or
Windows username. Stewardships - Documents lives in a different SharePoint
library than this project, so it can't be found by climbing relative
folders the way this project's own Master-data/Employee-data folders are.
Instead, find_onedrive_synced_folder() below looks for it under whoever's
home directory is running the script -- so this keeps working unchanged
after being copied to a new computer (e.g. when this moves off an intern's
laptop to its permanent home), as long as OneDrive is signed in and
syncing the Stewardships site there. See that function's docstring for
exactly how the search works and what to do if it ever can't find the
folder.
"""

import glob
import os
import warnings
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")  # openpyxl's noisy-but-harmless format warnings

# ============ CONFIG ============


def find_onedrive_synced_folder(folder_name, org_hint="gilroy"):
    """Locate a OneDrive-synced SharePoint folder (e.g. "Stewardships -
    Documents") without hardcoding a Windows username or computer name,
    so scripts keep working after being copied to a different computer
    or run under a different account.

    Search order:
      1. The exact layout confirmed to work today:
         <home>/GILROY KERNAN & GILROY/<folder_name>
      2. Any folder directly under <home> whose name contains `org_hint`
         (case-insensitive) -- covers OneDrive naming it slightly
         differently, e.g. "OneDrive - Gilroy Kernan & Gilroy" -- searched
         up to 3 levels deep for a directory named exactly `folder_name`.

    Returns the path as a string, or None if it can't be found anywhere.
    Set the matching environment variable (see where this is called
    below) to skip this search entirely and force a specific path.
    """
    home = Path.home()

    candidate = home / "GILROY KERNAN & GILROY" / folder_name
    if candidate.is_dir():
        return str(candidate)

    try:
        org_roots = [
            p for p in home.iterdir()
            if p.is_dir() and org_hint in p.name.lower()
        ]
    except OSError:
        org_roots = []

    for root in org_roots:
        for dirpath, dirnames, _ in os.walk(root):
            depth = len(Path(dirpath).relative_to(root).parts)
            if depth >= 3:
                dirnames[:] = []  # don't recurse further down this branch
            if os.path.basename(dirpath) == folder_name:
                return dirpath

    return None


# Stewardships - Documents lives in a different SharePoint library than this
# project, so it's located via find_onedrive_synced_folder() above rather
# than a relative path. Set the STEWARDSHIP_SOURCE_DIR environment variable
# to override/skip the search entirely (e.g. if it's ever synced somewhere
# find_onedrive_synced_folder() can't guess).
STEWARDSHIP_SOURCE_DIR = os.environ.get(
    "STEWARDSHIP_SOURCE_DIR"
) or find_onedrive_synced_folder("Stewardships - Documents")

# Client workbook filenames containing any of these (case-insensitive) are
# skipped -- they're sample/test documents, not real clients.
SKIP_NAME_MARKERS = ["sample", "test doc"]

# Timeline tab layout, verified identical across all 205 client workbooks:
# day-offset header row, actual calendar dates, then one row per tracked
# activity with milestone labels in the date columns.
TIMELINE_DATE_ROW = 5
TIMELINE_FIRST_ACTIVITY_ROW = 6
TIMELINE_FIRST_COL = 5   # column E (the "kickoff"/day-0 column)
TIMELINE_LAST_COL = 20   # column T (day 450)
TIMELINE_ACTIVITY_COL = 2   # column B
TIMELINE_OWNER_COL = 3      # column C

# ================================

ROOT = Path(__file__).resolve().parent.parent  # .../Code
PROJECT_ROOT = ROOT.parent                     # project root (parent of Code/)
MASTER_DIR = PROJECT_ROOT / "Master-data"
EMPLOYEE_DIR = PROJECT_ROOT / "Employee-data"
OUTPUT_FILENAME = "Whats_On_My_Plate_Stewardship_Knowledge.md"
OUTPUT = MASTER_DIR / OUTPUT_FILENAME
EMPLOYEE_OUTPUT = EMPLOYEE_DIR / OUTPUT_FILENAME

# Stewardship content has no Leadership-vs-everyone split the way Ninety
# rocks/Monday tasks do -- it's client account data, not internal work
# items -- so the exact same file goes to both Master-data and
# Employee-data rather than being filtered differently per audience.


def fmt_date(v):
    if not isinstance(v, datetime):
        return None
    return v.strftime("%B %-d, %Y") if os.name != "nt" else v.strftime("%B %#d, %Y")


def extract_cover(ws):
    client_name = ws["M10"].value
    meeting_date = fmt_date(ws["S14"].value)
    attendees = ws["S15"].value
    if isinstance(attendees, str):
        attendees = attendees.replace("\n", "; ").strip()
    return client_name, meeting_date, attendees


def extract_timeline(ws):
    """Returns a list of (date, activity, owner, milestone_label) tuples,
    sorted chronologically (undated entries sort last)."""
    col_to_date = {}
    for col in range(TIMELINE_FIRST_COL, TIMELINE_LAST_COL + 1):
        v = ws.cell(row=TIMELINE_DATE_ROW, column=col).value
        if isinstance(v, datetime):
            col_to_date[col] = v

    entries = []
    max_row = ws.max_row
    for row in range(TIMELINE_FIRST_ACTIVITY_ROW, max_row + 1):
        activity = ws.cell(row=row, column=TIMELINE_ACTIVITY_COL).value
        if activity is None or not str(activity).strip():
            continue
        owner = ws.cell(row=row, column=TIMELINE_OWNER_COL).value
        for col in range(TIMELINE_FIRST_COL, TIMELINE_LAST_COL + 1):
            label = ws.cell(row=row, column=col).value
            if label is None or not str(label).strip():
                continue
            entries.append((col_to_date.get(col), str(activity).strip(),
                             (owner or "").strip() if isinstance(owner, str) else owner,
                             str(label).strip()))

    entries.sort(key=lambda e: (e[0] is None, e[0]))
    return entries


def build_client_section(fname, client_name, meeting_date, attendees, timeline_entries, timeline_missing):
    display_name = client_name or Path(fname).stem
    lines = [f"## {display_name}", ""]
    lines.append(f"- Source file: `{fname}`")
    lines.append(f"- Most recent annual review meeting date: {meeting_date or 'not recorded'}")
    lines.append(f"- Attendees: {attendees or 'not recorded'}")

    if timeline_missing:
        lines.append("- Strategic timeline: no Timeline tab found in this workbook.")
    elif not timeline_entries:
        lines.append("- Strategic timeline: Timeline tab present but no milestones recorded.")
    else:
        lines.append("- Strategic timeline (all milestones, chronological):")
        for date, activity, owner, label in timeline_entries:
            date_str = date.strftime("%B %-d, %Y") if os.name != "nt" and isinstance(date, datetime) else (
                date.strftime("%B %#d, %Y") if isinstance(date, datetime) else "no date"
            )
            owner_str = f" (owner: {owner})" if owner else ""
            lines.append(f"  - {date_str}: {activity} — {label}{owner_str}")

    lines.append("")
    return "\n".join(lines)


def main():
    if not STEWARDSHIP_SOURCE_DIR or not os.path.isdir(STEWARDSHIP_SOURCE_DIR):
        raise SystemExit(
            "Couldn't find the Stewardships - Documents folder anywhere under "
            f"{Path.home()}.\n"
            "This usually means OneDrive isn't signed in / syncing the "
            "Stewardships site on this computer yet -- open OneDrive, make "
            "sure that site is syncing, and try again.\n"
            "If it's synced somewhere unusual, set the STEWARDSHIP_SOURCE_DIR "
            "environment variable to the exact path as a one-off override."
        )
    MASTER_DIR.mkdir(parents=True, exist_ok=True)
    EMPLOYEE_DIR.mkdir(parents=True, exist_ok=True)

    all_files = sorted(glob.glob(os.path.join(STEWARDSHIP_SOURCE_DIR, "*.xlsx")))
    files, skipped = [], []
    for f in all_files:
        base = os.path.basename(f).lower()
        if any(marker in base for marker in SKIP_NAME_MARKERS):
            skipped.append(os.path.basename(f))
        else:
            files.append(f)

    print(f"Found {len(all_files)} workbooks, skipping {len(skipped)} sample/test docs, processing {len(files)}.")

    sections = []
    no_cover, no_timeline, errors = [], [], []

    for path in files:
        fname = os.path.basename(path)
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as e:
            errors.append((fname, str(e)))
            continue

        if "Cover" in wb.sheetnames:
            client_name, meeting_date, attendees = extract_cover(wb["Cover"])
        else:
            no_cover.append(fname)
            client_name, meeting_date, attendees = None, None, None

        timeline_missing = "Timeline" not in wb.sheetnames
        timeline_entries = [] if timeline_missing else extract_timeline(wb["Timeline"])
        if timeline_missing:
            no_timeline.append(fname)

        wb.close()

        sections.append(build_client_section(
            fname, client_name, meeting_date, attendees, timeline_entries, timeline_missing
        ))

    now = datetime.now(timezone.utc)
    header = (
        "# Stewardship meeting attendees & strategic timelines\n\n"
        f"Data as of {now:%A, %B %-d, %Y} at {now:%H:%M} UTC. "
        if os.name != "nt" else
        f"Data as of {now:%A, %B %#d, %Y} at {now:%H:%M} UTC. "
    )
    header += (
        "This document lists, for each stewardship client, the most recent "
        "annual review meeting date and attendees from that client's Cover "
        "page, plus every milestone on their strategic timeline (from the "
        "Timeline tab) with its calendar date and owner (GKG, the client, "
        "or both).\n\n"
        f"Clients covered: {len(sections)}. "
        f"Missing a Cover tab: {len(no_cover)}. "
        f"Missing a Timeline tab: {len(no_timeline)}.\n\n"
    )

    content = header + "\n".join(sections)

    for out_path in (OUTPUT, EMPLOYEE_OUTPUT):
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(content)
        print(f"Wrote {out_path}")
    print(f"Clients: {len(sections)}  |  No Cover tab: {len(no_cover)}  |  "
          f"No Timeline tab: {len(no_timeline)}  |  Errors: {len(errors)}")
    for fname, msg in errors:
        print(f"  ERROR: {fname} -> {msg}")


if __name__ == "__main__":
    main()
